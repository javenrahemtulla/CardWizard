import glob
import io
import os
import re

import docx
import pytest
from docx.enum.text import WD_COLOR_INDEX

from app.ingest import parse_file, parse_html_string
from app.ingest.model import Para
from app.ingest.structure import extract_cards, parse_cite

BODY = ("The mitigation of extinction risk depends on institutional design and on the willingness of states to "
        "cooperate under uncertainty, which is why the evidence here matters for the round. ") * 4
CITE = "Smith 19 [Jane Smith, professor of political science at State University, 2019, “Extinction and Institutions,” Journal of Risk, https://example.org/a]"


def make_docx(with_h4=True):
    d = docx.Document()
    if with_h4:
        d.add_heading("Speech", 1)
        d.add_heading("Block", 3)
    tag = d.add_heading("Extinction risk is high", 4) if with_h4 else d.add_paragraph("Extinction risk is high")
    d.add_paragraph(CITE)
    p = d.add_paragraph()
    p.add_run("Plain start. ")
    r = p.add_run("underlined context ")
    r.underline = True
    r = p.add_run("spoken words")
    r.underline = True
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    p.add_run(" " + BODY)
    return d


def to_bytes(d):
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def tiers(card):
    out = {0: "", 1: "", 2: ""}
    for para in card.body:
        for t, tier in para:
            out[tier] += t
    return out


def test_highlight_is_spoken_underline_is_context():
    res = parse_file("a.docx", to_bytes(make_docx()))
    assert len(res.cards) == 1
    c = res.cards[0]
    t = tiers(c)
    assert t[2].strip() == "spoken words"
    assert "underlined context" in t[1]
    assert "Plain start" in t[0]
    assert c.tag == "Extinction risk is high"
    assert c.method == "formatted"


def test_unformatted_docx_finds_card_by_cite():
    res = parse_file("a.docx", to_bytes(make_docx(with_h4=False)))
    assert len(res.cards) == 1
    assert res.cards[0].method == "heuristic"
    assert res.cards[0].tag == "Extinction risk is high"


def test_plain_text():
    txt = f"Extinction risk is high\n{CITE}\n{BODY}\n"
    res = parse_file("a.txt", txt.encode())
    assert len(res.cards) == 1 and res.cards[0].tag == "Extinction risk is high"


def test_google_docs_style_html():
    html = f"""<html><head><style>.c1{{background-color:#ffff00}}.c2{{text-decoration:underline}}</style></head><body>
    <h4>Extinction risk is high</h4><p>{CITE}</p>
    <p>Plain <span class="c2">context</span> <span class="c1 c2">spoken here</span> {BODY}</p></body></html>"""
    res = parse_html_string(html)
    assert len(res.cards) == 1
    t = tiers(res.cards[0])
    assert t[2].strip() == "spoken here" and "context" in t[1]


def test_in_body_paragraph_starting_with_year_is_not_a_cite():
    paras = [Para([(s, 0, False)], 0) for s in (
        "Extinction risk is high", CITE, BODY,
        "In 2015, the union adopted its first package that should bring the economy toward circularity. " * 6,
        BODY)]
    cards, _ = extract_cards(paras)
    assert len(cards) == 1 and len(cards[0].body) == 3


def test_cite_and_body_in_one_paragraph_is_split():
    paras = [Para([("Extinction risk is high", 0, False)], 4),
             Para([(CITE + " " + BODY * 4, 0, False)], 0)]
    cards, _ = extract_cards(paras)
    assert len(cards) == 1
    assert cards[0].cite_text.endswith("https://example.org/a]")
    assert BODY[:30] in "".join(t for t, _ in cards[0].body[0])


def test_parse_cite():
    assert parse_cite(CITE) == ("Smith", 2019, "https://example.org/a")
    assert parse_cite("Pullen et al. 17 [Alison Pullen, 2017]")[:2] == ("Pullen et al.", 2017)


SAMPLES = sorted(glob.glob(os.path.join(os.path.dirname(__file__), "..", "samples", "*.docx")))


@pytest.mark.skipif(not SAMPLES, reason="no sample files in samples/")
def test_real_samples_recover_every_card_with_and_without_formatting():
    for f in SAMPLES:
        data = open(f, "rb").read()
        res = parse_file(f, data)
        assert res.cards, f
        # every card has spoken text, and spoken text is far smaller than the card
        for c in res.cards:
            t = tiers(c)
            assert t[2], (f, c.tag)
            assert len(t[2]) < 0.6 * sum(len(v) for v in t.values())
        from app.ingest.docx_reader import read_docx
        stripped = [Para([(r[0], 0, False) for r in p.runs], 0) for p in read_docx(data)]
        cards2, _ = extract_cards(stripped)
        assert {c.tag for c in cards2} == {c.tag for c in res.cards}, f
