import json
import os
import shutil
import subprocess

import pytest

from tests.test_ingest import make_docx, to_bytes
from tools import build_site

os.environ["CARDWIZARD_EMBEDDER"] = "hash"


def write(folder, name, doc):
    with open(os.path.join(folder, name), "wb") as f:
        f.write(to_bytes(doc))


def other_doc():
    import docx
    from docx.enum.text import WD_COLOR_INDEX
    d = docx.Document()
    d.add_heading("A completely different claim about trade", 4)
    d.add_paragraph("Jones 20 [Al Jones, professor of economics at Other University, 2020, “Trade Wars,” Journal of Trade, https://example.org/t]")
    p = d.add_paragraph("Tariffs raise prices. ")
    r = p.add_run("tariffs hurt consumers")
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    p.add_run(" and the effect compounds across supply chains over many years. " * 5)
    return d


def load(out):
    return json.load(open(os.path.join(out, "data", "index.json")))


def test_sync_add_dedupe_remove_and_idempotent(tmp_path):
    inc, out = str(tmp_path / "in"), str(tmp_path / "out")
    os.makedirs(inc)
    write(inc, "a.docx", make_docx())
    r = build_site.sync(inc, out)
    assert r["cards"] == 1 and r["files"] == 1 and r["vectors"] == 1
    idx = load(out)
    card = idx["cards"][0]
    assert card["spoken"].strip() == "spoken words" and card["tag"] == "Extinction risk is high"
    body = json.load(open(os.path.join(out, "data", "c", f"{card['id']}.json")))
    assert body["body"] and body["url"] == "https://example.org/a"

    # running again changes nothing
    before = open(os.path.join(out, "data", "index.json")).read()
    build_site.sync(inc, out)
    assert open(os.path.join(out, "data", "index.json")).read() == before

    # same card in a differently named copy with a different file hash: one card, two sources
    d = make_docx()
    d.paragraphs[0].insert_paragraph_before("A different speech title, so the file hash changes but the card does not", style="Heading 1")
    write(inc, "b.docx", d)
    write(inc, "c.docx", other_doc())
    r = build_site.sync(inc, out)
    idx = load(out)
    assert r["files"] == 3 and len(idx["cards"]) == 2
    shared = [c for c in idx["cards"] if len(c["src"]) == 2]
    assert len(shared) == 1

    # removing one of the two copies keeps the card, removing the last copy drops it
    os.remove(os.path.join(inc, "a.docx"))
    build_site.sync(inc, out)
    assert len(load(out)["cards"]) == 2
    os.remove(os.path.join(inc, "b.docx"))
    build_site.sync(inc, out)
    idx = load(out)
    assert [c["tag"] for c in idx["cards"]] == ["A completely different claim about trade"]
    assert sorted(os.listdir(os.path.join(out, "data", "c"))) == [f"{idx['cards'][0]['id']}.json"]
    ids = json.load(open(os.path.join(out, "data", "embeddings.ids.json")))
    assert ids == [idx["cards"][0]["id"]]
    assert os.path.getsize(os.path.join(out, "data", "embeddings.bin")) == 384


def test_unreadable_file_is_recorded_not_fatal(tmp_path):
    inc, out = str(tmp_path / "in"), str(tmp_path / "out")
    os.makedirs(inc)
    with open(os.path.join(inc, "broken.docx"), "wb") as f:
        f.write(b"PK not really a docx")
    write(inc, "good.docx", make_docx())
    r = build_site.sync(inc, out)
    assert r["cards"] == 1
    files = json.load(open(os.path.join(out, "data", "files.json")))
    assert any("error" in v for v in files.values())


@pytest.mark.skipif(shutil.which("node") is None, reason="node not installed")
def test_browser_search_engine_finds_cards(tmp_path):
    inc, out = str(tmp_path / "in"), str(tmp_path / "out")
    os.makedirs(inc)
    write(inc, "a.docx", make_docx())
    write(inc, "c.docx", other_doc())
    build_site.sync(inc, out)
    script = f"""
      const CS = require({json.dumps(os.path.abspath('docs/search.js'))});
      const idx = JSON.parse(require('fs').readFileSync({json.dumps(os.path.join(out, 'data', 'index.json'))}));
      const I = new CS.Index(idx.cards);
      const top = q => (I.keyword(q, 3).map(i => idx.cards[i].tag)[0] || null);
      console.log(JSON.stringify([top('tariffs hurt consumers'), top('extinction institutions'), top('nothing matches zzzz'), top('"tariffs hurt"')]));
    """
    res = subprocess.run(["node", "-e", script], capture_output=True, text=True, check=True)
    assert json.loads(res.stdout) == ["A completely different claim about trade", "Extinction risk is high", None,
                                      "A completely different claim about trade"]
